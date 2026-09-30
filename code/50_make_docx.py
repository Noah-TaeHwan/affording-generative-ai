"""
50_make_docx.py -- Word version of a manuscript (for readers or journals that need .docx).
Flattens the LaTeX source (inlines tables, resolves cross-references from a compiled .aux file, simplifies multi-row
table headers, renders the figure PDFs as PNG images) and converts it with pandoc (+citeproc, author-date references).
The PDF built by build_papers.sh remains the reference version; the Word file is a convenience copy.

Usage: python code/50_make_docx.py ko [output.docx]      (or en; default output/Affording_Generative_AI_<LANG>.docx)
Needs pandoc 3.1, python-docx, and TeX Live with poppler-utils (xelatex or pdflatex for the cross-references,
pdftoppm for the figures). All intermediate files are written to a temporary directory; the manuscript folder is
not modified.
"""
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

LANG = sys.argv[1] if len(sys.argv) > 1 else "ko"
ROOT = pathlib.Path(__file__).resolve().parents[1]
P = ROOT / f"paper_{LANG}"
OUT = pathlib.Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else ROOT / "output" / f"Affording_Generative_AI_{LANG.upper()}.docx"
TMP = tempfile.TemporaryDirectory()
BUILD = pathlib.Path(TMP.name)
src = (P / "main.tex").read_text(encoding="utf-8")

# ---------------------------------------------------------------- cross-references from a compiled .aux (in a copy)
work = BUILD / "tex"
shutil.copytree(P, work)
engine = "xelatex" if LANG == "ko" else "pdflatex"
for _ in range(2):
    subprocess.run([engine, "-interaction=nonstopmode", "-halt-on-error", "main.tex"], cwd=work,
                   capture_output=True, check=True)
aux = (work / "main.aux").read_text(encoding="utf-8")
labels = dict(re.findall(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", aux))

# ---------------------------------------------------------------- simplified single-row table headers
H = {
 "ko": {
  "tab:quintile": r"소득집단 & $N$ & 하위 20\% & 중위 20\% & 상위 20\% & 모바일 광대역 2GB (GNI 대비 \%) & AI 요금제/광대역 \\",
  "tab:local": r"요금제 & 미국 가격 & 고소득 & 상위중소득 & 하위중소득 & 저소득 & $\hat\beta$ & (표준오차) & $N$ \\",
  "tab:store": r"소득집단 & $N$ & 현지 스토어(\%) & 스토어 없음(\%) & 미제공(\%) & 달러 표시(\%) & 최저가 유료(달러) & 최저가 부담(\%) & Plus 부담(\%) \\",
  "tab:diffusion": r" & (1) 수준 & (2) 로그 & (3) 로그 & (4) 로그 & (5) 로그 & (6) 로그 & (7) 로그 \\",
  "tab:decomp": r"소득집단 & $N$ & AI 이용자(\%) & 인터넷(\%) & AI / 인터넷(\%) & 로그 격차 전체 & 연결성 & 비율 & 연결성 몫(\%) \\",
  "tab:auitime": r"시점 & 탄력성(전체) & (표준오차) & $N$ & $R^2$ & 탄력성(공통 표본) & (표준오차) \\",
 },
 "en": {
  "tab:quintile": r"Income group & $N$ & Poorest 20\% & Middle 20\% & Richest 20\% & Mobile broadband 2 GB (\% GNI p.c.) & AI plan / broadband \\",
  "tab:local": r"Plan & U.S. price & High & Upper-mid. & Lower-mid. & Low & $\hat\beta$ & (s.e.) & $N$ \\",
  "tab:store": r"Income group & $N$ & Local store (\%) & No store (\%) & Not offered (\%) & USD-priced (\%) & Cheapest tier (US\$) & Cheapest burden (\%) & Plus burden (\%) \\",
  "tab:diffusion": r" & (1) Level & (2) Log & (3) Log & (4) Log & (5) Log & (6) Log & (7) Log \\",
  "tab:decomp": r"Income group & $N$ & AI users (\%) & Internet (\%) & AI / internet (\%) & Log gap, total & Connectivity & Ratio & Connectivity share (\%) \\",
  "tab:auitime": r"Window & Elasticity (all) & (s.e.) & $N$ & $R^2$ & Elasticity (common) & (s.e.) \\",
 }}[LANG]
TABWORD, FIGWORD = ("표", "그림") if LANG == "ko" else ("Table", "Figure")
THEOREMS = ([("definition", "정의"), ("remark", "비고"), ("proposition", "명제")] if LANG == "ko"
            else [("definition", "Definition"), ("remark", "Remark"), ("proposition", "Proposition")])
PROOF = "증명." if LANG == "ko" else "Proof."


def inline_tables(t):
    return re.sub(r"\\tabinput\{(tables/[^}]+)\}", lambda m: (P / m.group(1)).read_text(encoding="utf-8"), t)


def multicol(t):
    # \multicolumn{n}{spec}{text} -> text followed by n-1 empty cells
    return re.sub(r"\\multicolumn\{(\d+)\}\{[^}]*\}\{((?:[^{}]|\{[^{}]*\})*)\}",
                  lambda m: m.group(2) + " &" * (int(m.group(1)) - 1), t)


def fix_table(env):
    lab = re.search(r"\\label\{(tab:[^}]+)\}", env)
    lab = lab.group(1) if lab else None
    notes = re.search(r"\\begin\{tablenotes\}(?:\[[^\]]*\])?(.*?)\\end\{tablenotes\}", env, re.S)
    note_txt = re.sub(r"\\(footnotesize|small)\b", "", notes.group(1)).replace("\\item", "").strip() if notes else ""
    env = re.sub(r"\\begin\{tablenotes\}.*?\\end\{tablenotes\}", "", env, flags=re.S)
    env = re.sub(r"\\(begin|end)\{threeparttable\}", "", env)
    env = re.sub(r"\\setlength\{\\tabcolsep\}\{[^}]*\}", "", env)
    env = re.sub(r"\\(footnotesize|scriptsize|small)\b", "", env)
    env = re.sub(r"\\cmidrule(\([a-z]*\))?\{[^}]*\}", "", env)
    env = env.replace("\\addlinespace", "")
    env = re.sub(r"\\begin\{tabularx\}\{[^}]*\}\{", r"\\begin{tabular}{", env).replace("\\end{tabularx}", "\\end{tabular}")

    def spec(m):
        s = m.group(1)
        s = re.sub(r">\{[^}]*\}", "", s)
        s = re.sub(r"p\{[^}]*\}", "l", s).replace("X", "l").replace("@{}", "")
        return "\\begin{tabular}{" + s + "}"
    env = re.sub(r"\\begin\{tabular\}\{((?:[^{}]|\{[^{}]*\})*)\}", spec, env)
    if lab in H:
        env = re.sub(r"(\\toprule)(.*?)(\\midrule)", lambda m: m.group(1) + "\n" + H[lab] + "\n" + m.group(3), env,
                     count=1, flags=re.S)
    env = multicol(env)
    num = labels.get(lab, "")
    env = re.sub(r"\\caption\{", lambda m: "\\caption{" + f"{TABWORD} {num}. ", env, count=1)
    return env + ("\n\n" + note_txt + "\n\n" if note_txt else "")


def png_figure(m):
    """Render the figure PDF used by the manuscript as a PNG in the build directory."""
    pdf_path = P / (m.group(1) + ".pdf")
    stem = BUILD / pathlib.Path(m.group(1)).name
    subprocess.run(["pdftoppm", "-png", "-r", "180", "-singlefile", str(pdf_path), str(stem)], check=True)
    return str(stem) + ".png"


def fix_figure(env):
    lab = re.search(r"\\label\{(fig:[^}]+)\}", env)
    num = labels.get(lab.group(1), "") if lab else ""
    env = re.sub(r"(figures/[A-Za-z0-9_]+)\.pdf", png_figure, env)
    env = re.sub(r"\\caption\{", lambda m: "\\caption{" + f"{FIGWORD} {num}. ", env, count=1)
    return env


# Materialise theorem numbering for Word; PDF labels remain authoritative.
for kind, title in THEOREMS:
    counter = [0]

    def theorem(m, counter=counter, title=title):
        counter[0] += 1
        return "\\paragraph{" + title + " " + str(counter[0]) + (" (" + m.group(1) + ")" if m.group(1) else "") + ".}"
    src = re.sub(r"\\begin\{" + kind + r"\}(?:\[([^]]+)\])?", theorem, src)
    src = src.replace("\\end{" + kind + "}", "")
src = src.replace("\\begin{proof}", "\\paragraph{" + PROOF + "}").replace("\\end{proof}", "")
t = inline_tables(src)
t = t.replace(r"$^\dagger$", "†")
t = re.sub(r"\\begin\{table\}.*?\\end\{table\}", lambda m: fix_table(m.group(0)), t, flags=re.S)
t = re.sub(r"\\begin\{figure\}.*?\\end\{figure\}", lambda m: fix_figure(m.group(0)), t, flags=re.S)


# equations: show their numbers explicitly
def eqnum(m):
    body = m.group(1)
    lab = re.search(r"\\label\{([^}]+)\}", body)
    n = labels.get(lab.group(1), "") if lab else ""
    body = re.sub(r"\\label\{[^}]+\}", "", body).strip()
    return "\\[\n" + body + (f" \\qquad ({n})" if n else "") + "\n\\]"


t = re.sub(r"\\begin\{equation\}(.*?)\\end\{equation\}", eqnum, t, flags=re.S)
t = re.sub(r"\\eqref\{([^}]+)\}", lambda m: f"({labels.get(m.group(1), '?')})", t)
t = re.sub(r"\\ref\{([^}]+)\}", lambda m: labels.get(m.group(1), "?"), t)
# section numbering as in the PDF (numbers, then appendix letters)
sec, sub, app = 0, 0, False
out = []
for line in t.split("\n"):
    if line.strip() == "\\appendix":
        app, sec = True, 0
        continue
    m = re.match(r"\\section\{(.*?)\}(.*)", line)
    if m:
        sec += 1
        sub = 0
        n = chr(64 + sec) if app else str(sec)
        line = f"\\section*{{{n}~~{m.group(1)}}}{m.group(2)}"
    m = re.match(r"\\subsection\{(.*?)\}(.*)", line)
    if m:
        sub += 1
        line = f"\\subsection*{{{sec}.{sub}~~{m.group(1)}}}{m.group(2)}"
    out.append(line)
t = "\n".join(out)
ref_title = "참고문헌" if LANG == "ko" else "References"
t = re.sub(r"\\bibliographystyle\{[^}]*\}\s*\\bibliography\{[^}]*\}",
           "\\\\section*{" + ref_title + "}\n\nREFSPLACEHOLDER\n", t)
t = t.replace("\\qed", "").replace("\\nolinkurl{", "\\texttt{").replace("\\clearpage", "").replace("\\newpage", "")
t = re.sub(r"\\FloatBarrier\b", "", t)
t = re.sub(r"\\makeatletter.*?\\makeatother[^\n]*", "", t)
(BUILD / "flat.tex").write_text(t, encoding="utf-8")

# ---------------------------------------------------------------- reference document with suitable fonts
ref = BUILD / "reference.docx"
with open(ref, "wb") as fh:
    subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], stdout=fh, check=True)
doc = Document(ref)
latin, east = ("Times New Roman", "Malgun Gothic") if LANG == "ko" else ("Times New Roman", "Times New Roman")
for st in doc.styles:
    try:
        rpr = st.element.get_or_add_rPr()
    except Exception:
        continue
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    for a in ["w:ascii", "w:hAnsi", "w:cs"]:
        fonts.set(qn(a), latin)
    fonts.set(qn("w:eastAsia"), east)
    for a in ["w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"]:
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]
for name, size in [("Normal", 10.5), ("Body Text", 10.5), ("First Paragraph", 10.5), ("Compact", 9)]:
    if name in [s.name for s in doc.styles]:
        doc.styles[name].font.size = Pt(size)
doc.save(ref)

lua = BUILD / "refs.lua"
lua.write_text('function Para(el)\n  if #el.content == 1 and el.content[1].t == "Str" and el.content[1].text == "REFSPLACEHOLDER" then\n'
               '    return pandoc.Div({}, pandoc.Attr("refs"))\n  end\nend\n')
OUT.parent.mkdir(parents=True, exist_ok=True)
abstract_title = "초록" if LANG == "ko" else "Abstract"
cmd = ["pandoc", str(BUILD / "flat.tex"), "-f", "latex", "-t", "docx", "--resource-path", str(BUILD) + ":" + str(P),
       "--reference-doc", str(ref), "--lua-filter", str(lua), "--citeproc",
       "-M", "abstract-title=" + abstract_title,
       "--bibliography", str(P / "references.bib"), "-o", str(OUT)]
r = subprocess.run(cmd, capture_output=True, text=True)
if r.returncode != 0:
    raise SystemExit(r.stderr[-3000:])
if r.stderr:
    print(r.stderr[-2000:])

# pandoc 3.1 ignores abstract-title for docx: relabel the abstract heading in place
d = Document(OUT)
for para in d.paragraphs[:12]:
    if para.text.strip() == "Abstract" and LANG == "ko":
        for run in para.runs:
            run.text = ""
        para.runs[0].text = abstract_title
        break
d.save(OUT)
print("written", OUT)
